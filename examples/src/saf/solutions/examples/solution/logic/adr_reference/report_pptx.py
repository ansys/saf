# Copyright (C) 2026 ANSYS, Inc. and/or its affiliates.
# SPDX-License-Identifier: Apache-2.0

#
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

# ©2026, ANSYS Inc. Unauthorized use, distribution or duplication is prohibited.
"""Input PowerPoint deck of the ADR reference report.

ADR builds PowerPoint reports from an input deck: every placeholder whose name starts with
``Nexus:`` is filled with the items returned by the query that follows, for example
``Nexus:query=A|i_name|eq|my-item;``. The placeholder type must match the item type:

- String and HTML items go into text (body) placeholders.
- Table and Tree items go into table placeholders.
- Image, Scene, and File items go into picture placeholders. Scenes and files are only shown
  when ADR has generated a preview picture for them.

The deck is generated with python-pptx (a dependency of ansys-dynamicreporting-core) and stored in
the ADR database as a File item. You can also design your own deck in PowerPoint and rename its
placeholders in the selection pane.
"""

import logging
from pathlib import Path
import tempfile

from ansys.dynamicreporting.core.serverless import ADR, File, Item
from pptx import Presentation
from pptx.slide import Slide

from saf.solutions.examples.solution.logic.adr_reference.report_definition import (
    PPTX_INPUT_ITEM_NAME,
    REPORT_SECTIONS,
    REPORT_TITLE,
    ReportSection,
)

logger = logging.getLogger(__name__)

_TITLE_SLIDE_LAYOUT_INDEX = 0
_CONTENT_WITH_CAPTION_LAYOUT_INDEX = 7
_CONTENT_PLACEHOLDER_INDEX = 1
_CAPTION_PLACEHOLDER_INDEX = 2

# OOXML placeholder types accepted by ADR for each item type.
_PLACEHOLDER_TYPE_BY_ITEM_TYPE = {
    "string": "body",
    "html": "body",
    "table": "tbl",
    "tree": "tbl",
    "image": "pic",
    "scene": "pic",
    "file": "pic",
}


def build_pptx_input_deck(deck_path: Path) -> Path:
    """Build the input deck: a title slide, then one slide per section of the report.

    Parameters
    ----------
    deck_path : Path
        Path of the PowerPoint file to create.

    Returns
    -------
    Path
        Path of the created PowerPoint file.
    """
    presentation = Presentation()

    title_slide = presentation.slides.add_slide(presentation.slide_layouts[_TITLE_SLIDE_LAYOUT_INDEX])
    title_slide.shapes.title.text = REPORT_TITLE
    title_slide.placeholders[1].text = "Generated with serverless Ansys Dynamic Reporting"

    for section in REPORT_SECTIONS:
        slide = presentation.slides.add_slide(presentation.slide_layouts[_CONTENT_WITH_CAPTION_LAYOUT_INDEX])
        _fill_section_slide(slide, section)

    deck_path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(deck_path))
    return deck_path


def _fill_section_slide(slide: Slide, section: ReportSection) -> None:
    """Write the section title and description, and turn the content placeholder into an ADR placeholder."""
    slide.shapes.title.text = section.title
    slide.placeholders[_CAPTION_PLACEHOLDER_INDEX].text = section.description

    content_placeholder = slide.placeholders[_CONTENT_PLACEHOLDER_INDEX]
    content_placeholder.element.ph.set("type", _PLACEHOLDER_TYPE_BY_ITEM_TYPE[section.item_type])
    content_placeholder.name = f"Nexus:query=A|i_name|eq|{section.item_name};"


def delete_pptx_input_item(adr: ADR) -> None:
    """Delete the input deck from the ADR database, together with its media file."""
    items = adr.query(File, query=f"A|i_name|eq|{PPTX_INPUT_ITEM_NAME};")
    if items:
        items.delete()


def create_pptx_input_item(adr: ADR, deck_path: Path) -> Item:
    """Store the input deck as the File item referenced by the PPTX layout template."""
    logger.info("Storing the input PowerPoint deck of the ADR reference report.")
    return adr.create_item(File, name=PPTX_INPUT_ITEM_NAME, content=str(deck_path), tags="pptx_input")


def recreate_pptx_input_item(adr: ADR) -> None:
    """Build the input deck and store it in the ADR database, replacing the previous one.

    The deck is shared by all projects. It is rebuilt at every report generation, so that it always
    matches this code.
    """
    delete_pptx_input_item(adr)
    with tempfile.TemporaryDirectory() as scratch_directory:
        deck_path = build_pptx_input_deck(Path(scratch_directory) / "adr_reference_input.pptx")
        create_pptx_input_item(adr, deck_path)
