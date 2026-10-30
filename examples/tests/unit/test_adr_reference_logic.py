# Copyright (C) 2026 ANSYS, Inc. and/or its affiliates.
# SPDX-License-Identifier: Apache-2.0

#
#
# Licensed under the Apache License, Version 2.0 (the "License");
# you may not use this file except in compliance with the License.
# You may obtain a copy of the License at
#
#     http://www.apache.org/licenses/LICENSE-2.0
#
# Unless required by applicable law or agreed to in writing, software
# distributed under the License is distributed on an "AS IS" BASIS,
# WITHOUT WARRANTIES OR CONDITIONS OF ANY KIND, either express or implied.
# See the License for the specific language governing permissions and
# limitations under the License.

"""Unit tests for the logic of the ADR reference example (no ADR installation needed)."""

import io
from pathlib import Path
from types import SimpleNamespace
from typing import Any
from unittest.mock import MagicMock
import zipfile

from ansys.dynamicreporting.core.serverless import (
    HTML,
    File,
    Image,
    PanelLayout,
    PPTXLayout,
    Scene,
    String,
    Table,
    Tree,
)
import flask
from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
import pytest

from saf.solutions.examples.solution.logic.adr_reference import (
    report_exports,
    report_items,
    report_pptx,
    report_templates,
)
from saf.solutions.examples.solution.logic.adr_reference.adr_settings import (
    AdrNotAvailableError,
    check_adr_availability,
    get_adr_paths,
)
from saf.solutions.examples.solution.logic.adr_reference.report_definition import (
    PPTX_INPUT_ITEM_NAME,
    PPTX_TEMPLATE_NAME,
    REPORT_SECTIONS,
    REPORT_TEMPLATE_NAME,
    build_item_filter,
    build_item_tags,
    build_project_tag,
)
from saf.solutions.examples.solution.logic.adr_reference.report_routes import register_report_asset_routes

PROJECT_TAG = "project=1234"


@pytest.fixture
def adr_environment(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Path:
    """Configure a valid ADR environment pointing to temporary directories."""
    installation_directory = tmp_path / "ADR"
    installation_directory.mkdir()
    monkeypatch.setenv("ADR_INSTALLATION_DIRECTORY", str(installation_directory))
    monkeypatch.setenv("GLOW_PROJECT_FILES_DIRECTORY", str(tmp_path / "glow" / "project_files"))
    return tmp_path


# ------------------------------------------------------------------------------------------------ #
# Settings
# ------------------------------------------------------------------------------------------------ #


def test_adr_is_not_available_when_installation_directory_is_not_set(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify ADR is reported as unavailable without ADR_INSTALLATION_DIRECTORY."""
    monkeypatch.setenv("ADR_INSTALLATION_DIRECTORY", "")

    availability = check_adr_availability()

    assert not availability.is_available
    assert "ADR_INSTALLATION_DIRECTORY" in availability.message


def test_adr_is_not_available_when_installation_directory_does_not_exist(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """Verify ADR is reported as unavailable when the installation directory is missing."""
    monkeypatch.setenv("ADR_INSTALLATION_DIRECTORY", str(tmp_path / "missing"))
    monkeypatch.setenv("GLOW_PROJECT_FILES_DIRECTORY", str(tmp_path))

    availability = check_adr_availability()

    assert not availability.is_available
    assert "does not exist" in availability.message


def test_adr_is_not_available_when_project_files_directory_is_not_set(
    adr_environment: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Verify ADR is reported as unavailable without GLOW_PROJECT_FILES_DIRECTORY."""
    monkeypatch.setenv("GLOW_PROJECT_FILES_DIRECTORY", "")

    availability = check_adr_availability()

    assert not availability.is_available
    assert "GLOW_PROJECT_FILES_DIRECTORY" in availability.message


def test_adr_paths_are_located_next_to_project_files(adr_environment: Path) -> None:
    """Verify the ADR directories are derived from the project files directory."""
    assert check_adr_availability().is_available

    paths = get_adr_paths()

    assert paths.installation_directory == adr_environment / "ADR"
    assert paths.database_directory == adr_environment / "glow" / "adr" / "report_data"
    assert paths.media_directory == adr_environment / "glow" / "adr" / "report_data" / "media"
    assert paths.static_directory == adr_environment / "glow" / "adr" / "static_assets"


def test_get_adr_paths_raises_when_adr_is_not_available(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify an explicit error is raised when ADR is not configured."""
    monkeypatch.setenv("ADR_INSTALLATION_DIRECTORY", "")

    with pytest.raises(AdrNotAvailableError, match="ADR_INSTALLATION_DIRECTORY"):
        get_adr_paths()


# ------------------------------------------------------------------------------------------------ #
# Report definition
# ------------------------------------------------------------------------------------------------ #


def test_report_sections_cover_every_common_item_type_except_animation() -> None:
    """Verify the report has one uniquely named section per common item type, without animation."""
    item_types = [section.item_type for section in REPORT_SECTIONS]

    assert item_types == ["string", "html", "table", "tree", "image", "scene", "file"]
    assert len({section.item_name for section in REPORT_SECTIONS}) == len(REPORT_SECTIONS)
    assert len({section.template_name for section in REPORT_SECTIONS}) == len(REPORT_SECTIONS)
    assert all(section.title and section.description for section in REPORT_SECTIONS)


def test_item_filter_selects_report_items_of_one_project() -> None:
    """Verify the tags and filter isolate the report items of a project."""
    project_tag = build_project_tag("1234")

    assert project_tag == PROJECT_TAG
    assert build_item_tags(project_tag) == "adr_reference project=1234"
    assert build_item_filter(project_tag) == "A|i_tags|cont|adr_reference;A|i_tags|cont|project=1234;"


# ------------------------------------------------------------------------------------------------ #
# Items
# ------------------------------------------------------------------------------------------------ #


def test_create_report_items_creates_one_item_per_section(tmp_path: Path) -> None:
    """Verify every section gets its example item, with the expected type, name, and tags."""
    adr = MagicMock()
    files = report_items.ReportItemFiles(
        image=tmp_path / "image.png", scene=tmp_path / "scene.avz", file=tmp_path / "file.txt"
    )

    report_items.create_report_items(adr, PROJECT_TAG, files)

    created = [(call.args[0], call.kwargs["name"], call.kwargs["tags"]) for call in adr.create_item.call_args_list]
    expected_types = [String, HTML, Table, Tree, Image, Scene, File]
    assert created == [
        (item_type, section.item_name, "adr_reference project=1234")
        for item_type, section in zip(expected_types, REPORT_SECTIONS, strict=True)
    ]
    file_contents = [call.kwargs["content"] for call in adr.create_item.call_args_list[4:]]
    assert file_contents == [str(files.image), str(files.scene), str(files.file)]


def test_table_item_has_labels_and_is_saved() -> None:
    """Verify the table item gets row and column labels before being saved again."""
    adr = MagicMock()

    table_item = report_items.create_table_item(adr, "tags")

    content = adr.create_item.call_args.kwargs["content"]
    assert content.shape == (5, 3)
    assert table_item.labels_column == ["Time [s]", "Displacement [mm]", "Stress [MPa]"]
    assert len(table_item.labels_row) == 5
    table_item.save.assert_called_once()


def test_tree_item_content_is_valid() -> None:
    """Verify the tree item content follows the format expected by ADR."""
    adr = MagicMock()

    report_items.create_tree_item(adr, "tags")

    def check_nodes(nodes: list[dict[str, Any]]) -> None:
        for node in nodes:
            assert {"key", "name", "value"} <= node.keys()
            check_nodes(node.get("children", []))

    check_nodes(adr.create_item.call_args.kwargs["content"])


def test_html_item_only_uses_tags_supported_by_pptx_export() -> None:
    """Verify the HTML item only uses the tags supported by the PowerPoint export."""
    adr = MagicMock()

    report_items.create_html_item(adr, "tags")

    content = adr.create_item.call_args.kwargs["content"]
    for unsupported_tag in ("<p", "<ul", "<li", "<b>", "<i>", "<div", "<span"):
        assert unsupported_tag not in content


def test_delete_report_items_only_deletes_items_of_the_project() -> None:
    """Verify existing report items are queried with the project filter and deleted."""
    adr = MagicMock()

    report_items.delete_report_items(adr, PROJECT_TAG)

    assert adr.query.call_args.kwargs["query"] == build_item_filter(PROJECT_TAG)
    adr.query.return_value.delete.assert_called_once()


def test_delete_report_items_does_nothing_without_items() -> None:
    """Verify nothing is deleted when no report item exists."""
    adr = MagicMock()
    adr.query.return_value = []

    report_items.delete_report_items(adr, PROJECT_TAG)


# ------------------------------------------------------------------------------------------------ #
# Templates
# ------------------------------------------------------------------------------------------------ #


def test_report_templates_are_recreated() -> None:
    """Verify existing templates are deleted, then the HTML and PowerPoint report templates are created."""
    adr = MagicMock()
    existing_templates = MagicMock()
    adr.query.return_value = existing_templates

    report_templates.recreate_report_templates(adr)

    deleted_names = [call.kwargs["query"] for call in adr.query.call_args_list]
    assert deleted_names == [f"A|t_name|eq|{REPORT_TEMPLATE_NAME};", f"A|t_name|eq|{PPTX_TEMPLATE_NAME};"]
    assert existing_templates.delete.call_count == 2

    created = [(call.args[0], call.kwargs["name"]) for call in adr.create_template.call_args_list]
    assert created[0][1] == REPORT_TEMPLATE_NAME
    section_templates = [name for template_type, name in created if template_type is PanelLayout]
    assert section_templates == [section.title for section in REPORT_SECTIONS]
    assert (PPTXLayout, PPTX_TEMPLATE_NAME) in created

    pptx_template = adr.create_template.return_value
    assert pptx_template.input_pptx == PPTX_INPUT_ITEM_NAME
    assert pptx_template.use_all_slides == "1"


def test_section_template_shows_explanation_and_selects_its_item() -> None:
    """Verify a section template has a header with the explanation and a filter on its example item."""
    adr = MagicMock()
    section = REPORT_SECTIONS[0]

    section_template = report_templates.create_section_template(adr, MagicMock(), section, 1)

    section_template.set_html.assert_called_once_with(f"<h2>1. {section.title}</h2><p>{section.description}</p>")
    section_template.set_filter.assert_called_once_with(f"A|i_name|eq|{section.item_name};")
    section_template.save.assert_called_once()


def test_partially_created_report_template_is_deleted() -> None:
    """Verify a failure while creating the sections deletes the root template and is re-raised."""
    adr = MagicMock()
    root_template = MagicMock()
    adr.create_template.side_effect = [root_template, RuntimeError("Template creation failed")]

    with pytest.raises(RuntimeError, match="Template creation failed"):
        report_templates.create_report_template(adr)

    root_template.delete.assert_called_once()


# ------------------------------------------------------------------------------------------------ #
# PowerPoint input deck
# ------------------------------------------------------------------------------------------------ #


def test_pptx_input_deck_has_one_adr_placeholder_per_section(tmp_path: Path) -> None:
    """Verify each section slide has a placeholder named with an ADR query and a matching type."""
    deck_path = report_pptx.build_pptx_input_deck(tmp_path / "deck" / "input.pptx")

    expected_types = {
        "string": PP_PLACEHOLDER.BODY,
        "html": PP_PLACEHOLDER.BODY,
        "table": PP_PLACEHOLDER.TABLE,
        "tree": PP_PLACEHOLDER.TABLE,
        "image": PP_PLACEHOLDER.PICTURE,
        "scene": PP_PLACEHOLDER.PICTURE,
        "file": PP_PLACEHOLDER.PICTURE,
    }
    slides = list(Presentation(str(deck_path)).slides)
    assert len(slides) == len(REPORT_SECTIONS) + 1
    for slide, section in zip(slides[1:], REPORT_SECTIONS, strict=True):
        adr_placeholders = [shape for shape in slide.placeholders if shape.name.startswith("Nexus:")]
        assert len(adr_placeholders) == 1
        assert adr_placeholders[0].name == f"Nexus:query=A|i_name|eq|{section.item_name};"
        assert adr_placeholders[0].placeholder_format.type == expected_types[section.item_type]
        assert slide.shapes.title.text == section.title


def test_pptx_input_deck_is_rebuilt_and_stored_as_a_file_item() -> None:
    """Verify the previous input deck is deleted, then the deck is rebuilt and stored as a File item."""
    adr = MagicMock()
    previous_deck = MagicMock()
    adr.query.return_value = previous_deck
    stored_slide_counts = []

    def create_item(item_type: type, **kwargs: Any) -> MagicMock:
        stored_slide_counts.append(len(Presentation(kwargs["content"]).slides))
        return MagicMock()

    adr.create_item.side_effect = create_item

    report_pptx.recreate_pptx_input_item(adr)

    previous_deck.delete.assert_called_once()
    assert adr.create_item.call_args.args[0] is File
    assert adr.create_item.call_args.kwargs["name"] == PPTX_INPUT_ITEM_NAME
    assert stored_slide_counts == [len(REPORT_SECTIONS) + 1]


# ------------------------------------------------------------------------------------------------ #
# Exports
# ------------------------------------------------------------------------------------------------ #


def test_html_archive_contains_the_exported_report_and_its_assets() -> None:
    """Verify the HTML export directory, including asset folders, is zipped."""

    def export_report_as_html(output_directory: Path, **kwargs: Any) -> Path:
        (output_directory / kwargs["filename"]).write_text("<html></html>")
        (output_directory / "media").mkdir()
        (output_directory / "media" / "image.png").write_bytes(b"png")
        return output_directory / kwargs["filename"]

    adr = MagicMock()
    adr.export_report_as_html.side_effect = export_report_as_html

    archive_content = report_exports.export_report_html_archive(adr, "filter")

    with zipfile.ZipFile(io.BytesIO(archive_content)) as archive:
        assert sorted(archive.namelist()) == [report_exports.HTML_EXPORT_FILE_NAME, "media/image.png"]
    assert adr.export_report_as_html.call_args.kwargs["name"] == REPORT_TEMPLATE_NAME
    assert adr.export_report_as_html.call_args.kwargs["item_filter"] == "filter"


def test_pptx_export_uses_the_pptx_template() -> None:
    """Verify the PowerPoint export renders the PPTX template."""
    adr = MagicMock()
    adr.render_report_as_pptx.return_value = b"pptx"

    assert report_exports.export_report_pptx(adr, "filter") == b"pptx"
    adr.render_report_as_pptx.assert_called_once_with(name=PPTX_TEMPLATE_NAME, item_filter="filter")


def test_pdf_export_uses_browser_rendering_when_available() -> None:
    """Verify the browser-based PDF rendering of recent ADR versions is preferred."""
    adr = MagicMock()
    adr.render_report_as_browser_pdf.return_value = b"pdf"

    assert report_exports.export_report_pdf(adr, "filter") == b"pdf"
    adr.render_report_as_pdf.assert_not_called()


def test_pdf_export_falls_back_to_weasyprint_rendering() -> None:
    """Verify older ADR versions render PDF files with render_report_as_pdf."""
    adr = MagicMock(spec=["render_report_as_pdf"])
    adr.render_report_as_pdf.return_value = b"pdf"

    assert report_exports.export_report_pdf(adr, "filter") == b"pdf"
    adr.render_report_as_pdf.assert_called_once_with(name=REPORT_TEMPLATE_NAME, item_filter="filter")


# ------------------------------------------------------------------------------------------------ #
# Routes
# ------------------------------------------------------------------------------------------------ #


def _dash_app_stub() -> SimpleNamespace:
    """Create an object exposing a Flask server, like a Dash application."""
    return SimpleNamespace(server=flask.Flask(__name__))


def test_report_asset_routes_serve_adr_files(adr_environment: Path) -> None:
    """Verify the media and static files of ADR are served by the registered routes."""
    paths = get_adr_paths()
    paths.media_directory.mkdir(parents=True)
    paths.static_directory.mkdir(parents=True)
    (paths.media_directory / "image.png").write_bytes(b"image")
    (paths.static_directory / "style.css").write_text("body {}")
    app = _dash_app_stub()

    register_report_asset_routes(app)

    client = app.server.test_client()
    assert client.get("/media/image.png").data == b"image"
    assert client.get("/static_assets/style.css").data == b"body {}"
    assert client.get("/media/missing.png").status_code == 404


def test_report_asset_routes_return_404_when_adr_is_not_available(monkeypatch: pytest.MonkeyPatch) -> None:
    """Verify the routes do not fail when ADR is not configured."""
    monkeypatch.setenv("ADR_INSTALLATION_DIRECTORY", "")
    app = _dash_app_stub()

    register_report_asset_routes(app)

    assert app.server.test_client().get("/media/image.png").status_code == 404


def test_report_asset_routes_keep_existing_routes() -> None:
    """Verify routes already registered by another report are not registered twice."""
    app = _dash_app_stub()
    app.server.add_url_rule("/media/<path:filename>", endpoint="existing_media", view_func=lambda filename: "")

    register_report_asset_routes(app)

    media_rules = [rule for rule in app.server.url_map.iter_rules() if rule.rule == "/media/<path:filename>"]
    assert [rule.endpoint for rule in media_rules] == ["existing_media"]
